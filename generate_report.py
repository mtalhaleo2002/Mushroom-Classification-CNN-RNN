
from pathlib import Path
import subprocess
import sys
import textwrap

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt

BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "outputs"
SHOTS = BASE / "screenshots"
REPORTS = BASE / "reports"

for folder in (OUTPUT, SHOTS, REPORTS):
    folder.mkdir(exist_ok=True)

FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"


def render_text(content, prefix, title):
    """Create readable code/output PNG pages."""
    font = ImageFont.truetype(FONT_PATH, 18)
    header_font = ImageFont.truetype(FONT_PATH, 23)

    lines = []
    for line in content.splitlines():
        lines.extend(
            textwrap.wrap(
                line,
                width=95,
                replace_whitespace=False,
                drop_whitespace=False
            ) or [""]
        )

    paths = []

    for start in range(0, len(lines), 32):
        part = start // 32 + 1
        chunk = lines[start:start + 32]

        image = Image.new(
            "RGB",
            (1500, 110 + len(chunk) * 31),
            "#1e1e1e"
        )

        draw = ImageDraw.Draw(image)
        draw.rectangle(
            (0, 0, 1500, 65),
            fill="#252526"
        )

        draw.text(
            (30, 18),
            f"{title} - Part {part}",
            font=header_font,
            fill="white"
        )

        for i, line in enumerate(chunk):
            draw.text(
                (30, 85 + i * 31),
                line,
                font=font,
                fill="#d4d4d4"
            )

        path = SHOTS / f"{prefix}_part{part}.png"
        image.save(path)
        paths.append(path)

    return paths


def add_images(doc, paths):
    for path in paths:
        doc.add_picture(str(path), width=Inches(6.7))


def read_source(filename):
    path = BASE / filename
    if not path.exists():
        raise FileNotFoundError(filename)
    return path.read_text(encoding="utf-8")


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(
        "Artificial Intelligence Using Python",
        0
    )
    doc.add_paragraph(
        "DigiSkills 3.0 - Batch 04\n"
        "Hands-on Exercise No. 04\n"
        "CNN, RNN and Generative AI Tokenization"
    )

    # -----------------------------
    # TASK 1: DEEP LEARNING
    # -----------------------------

    doc.add_heading(
        "Task 1: CNN and RNN Classification",
        1
    )

    doc.add_paragraph(
        "The Mushroom dataset was preprocessed using "
        "categorical encoding and split into 70% training "
        "and 30% testing data. CNN and RNN models were "
        "trained using TensorFlow/Keras."
    )

    source = read_source("task1_deep_learning.py")

    doc.add_heading("Python Code", 2)
    add_images(
        doc,
        render_text(
            source,
            "task1_code",
            "CNN and RNN Python Code"
        )
    )

    comparison_path = OUTPUT / "task1_model_comparison.csv"

    if not comparison_path.exists():
        raise FileNotFoundError(
            "Run task1_deep_learning.py first!"
        )

    comparison = pd.read_csv(comparison_path)

    doc.add_heading("Model Evaluation Results", 2)

    doc.add_paragraph(comparison.to_string(index=False))

    add_images(
        doc,
        render_text(
            comparison.to_string(index=False),
            "task1_results",
            "CNN vs RNN Evaluation Results"
        )
    )

    if comparison["Test_Accuracy"].notna().all():
        max_accuracy = comparison["Test_Accuracy"].max()
        winners = comparison.loc[
            comparison["Test_Accuracy"] == max_accuracy,
            "Model"
        ].tolist()

        if len(winners) == 1:
            conclusion = (
                f"{winners[0]} achieved the higher "
                f"test accuracy of {max_accuracy:.2%}."
            )
        else:
            conclusion = (
                "CNN and RNN achieved equal test accuracy "
                f"of {max_accuracy:.2%}."
            )

        doc.add_paragraph(conclusion)

    doc.add_heading("Training and Evaluation Graphs", 2)

    graphs = [
        "task1_accuracy_comparison.png",
        "task1_cnn_training.png",
        "task1_rnn_training.png",
        "task1_cnn_confusion.png",
        "task1_rnn_confusion.png"
    ]

    for graph_name in graphs:
        path = OUTPUT / graph_name
        if not path.exists():
            raise FileNotFoundError(
                f"Required graph missing: {graph_name}"
            )
        doc.add_paragraph(graph_name)
        doc.add_picture(str(path), width=Inches(6.3))

    doc.add_paragraph(
        "Note: The feature ordering in this tabular dataset "
        "does not represent a natural time sequence. "
        "The CNN and RNN models are experimental "
        "architectures for the assignment."
    )

    # -----------------------------
    # TASK 2: TOKENIZATION
    # -----------------------------

    doc.add_page_break()

    doc.add_heading(
        "Task 2: Generative AI Text Tokenization",
        1
    )

    doc.add_paragraph(
        "The exercise_text.json file was loaded and "
        "processed using a Hugging Face Transformers "
        "tokenizer. Text was converted into tokens, "
        "token IDs, and attention masks."
    )

    source = read_source("task2_tokenization.py")

    doc.add_heading("Python Code", 2)

    add_images(
        doc,
        render_text(
            source,
            "task2_code",
            "Hugging Face Tokenization Code"
        )
    )

    result = subprocess.run(
        [sys.executable, str(BASE / "task2_tokenization.py")],
        cwd=BASE,
        capture_output=True,
        text=True,
        timeout=180
    )

    if result.returncode != 0:
        print(result.stderr)
        raise RuntimeError("Task 2 execution failed!")

    doc.add_heading("Execution Output", 2)

    add_images(
        doc,
        render_text(
            result.stdout,
            "task2_output",
            "Tokenization Terminal Output"
        )
    )

    doc.add_heading("Conclusion", 2)

    doc.add_paragraph(
        "The CNN and RNN models were trained and "
        "evaluated on the Mushroom classification "
        "dataset. Their test accuracies were compared "
        "and training visualizations were generated. "
        "Text tokenization was also performed using "
        "a Hugging Face Transformers tokenizer."
    )

    report_path = REPORTS / "AIP301_Exercise_04_Solution.docx"

    doc.save(report_path)

    print("\nREPORT GENERATED SUCCESSFULLY!")
    print("Word File:", report_path)


if __name__ == "__main__":
    main()
